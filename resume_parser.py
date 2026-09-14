import os
from pathlib import Path
import time
from urllib import response
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel

load_dotenv()
my_api_key = os.getenv("GROQ_API_KEY")
if not my_api_key:
    raise ValueError("GROQ_API_KEY environment variable is not set.")
client=Groq(api_key=my_api_key)
model="openai/gpt-oss-20b"

job_description = """Description
Amazon Leo is Amazon's low Earth orbit satellite network. Our mission is to deliver fast, reliable internet connectivity to customers beyond the reach of existing networks. From individual households to schools, hospitals, businesses, and government agencies, Amazon Leo will serve people and organizations operating in locations without reliable connectivity.

Export Control Requirement: Due to applicable export control laws and regulations, candidates must be a U.S. citizen or national, U.S. permanent resident (i.e., current Green Card holder), or lawfully admitted into the U.S. as a refugee or granted asylum.

Key job responsibilities
Have you wanted an opportunity to develop the core wireless and networking for an advanced global telecom service spanning both space and terrestrial networks? Have you wondered what it takes to solve a multi-dimentional routing problem at global scale? Amazon Leo's SDN team will design, implement and operate Leo's network control plane. This service will enable a high throughput telecom service using a system of LEO satellites, customer terminals, gateways, cloud services, and terrestrial network infrastructure that connect into public and private networks. For this role, you will:
- Take responsibility for designing and delivering a modern, flexible, scalable, high throughput, low latency network architecture
- Implement Routing and Forwarding, Traffic Engineering, and SDN technologies to deliver the best-in-class broadband network services for our customers
- Implement highly-available software that will drive networking control plane operation on physical and remote nodes
- Define and automate processes used for software development, implementation, testing, and maintenance
- Provide operational excellence for control plane services, to including enabling tools and definition/refinement of operational processes
- Developing routing control plane software for network nodes in C and Rust

A day in the life
This is an opportunity to be at the heart of the design, building and operation of a massive satellite constellation and terrestrial network. We interact with the satellite ground control services, the customer engagement systems and monitoring services to constantly fine-tune our network and deliver a high quality service to our customers. You will be responsible to build the best network for our customers.

About the team
The SDN Control Plane team is building the brain driving the Leo Network. Innovation is in our DNA and we tackle the problems presented by a dynamic network at global scale. It is our responsibility to ensure that the network works and our customer has a one-of-a-kind experience using our service.

Our team will design, implement and operate the control plane that provides a high throughput telecom service comprised of Low Earth Orbit satellites, customer terminals, gateways, Cloud services and terrestrial network infrastructure that connects into public and private networks.

Basic Qualifications
- 3+ years of non-internship professional software development experience
- 2+ years of non-internship design or architecture (design patterns, reliability and scaling) of new and existing systems experience
- Experience programming with at least one software programming language

Preferred Qualifications
- 3+ years of full software development life cycle, including coding standards, code reviews, source control management, build processes, testing, and operations experience
- Bachelor's degree in computer science or equivalent

Amazon is an equal opportunity employer and does not discriminate on the basis of protected veteran status, disability, or other legally protected status.

Our inclusive culture empowers Amazonians to deliver the best results for our customers. If you have a disability and need a workplace accommodation or adjustment during the application and hiring process, including support for the interview or onboarding process, please visit https://amazon.jobs/content/en/how-we-hire/accommodations for more information. If the country/region you’re applying in isn’t listed, please contact your Recruiting Partner.


The base salary range for this position is listed below. Your Amazon package will include sign-on payments and restricted stock units (RSUs). Final compensation will be determined based on factors including experience, qualifications, and location. Amazon also offers comprehensive benefits including health insurance (medical, dental, vision, prescription, Basic Life & AD&D insurance and option for Supplemental life plans, EAP, Mental Health Support, Medical Advice Line, Flexible Spending Accounts, Adoption and Surrogacy Reimbursement coverage), 401(k) matching, paid time off, and parental leave. Learn more about our benefits at https://amazon.jobs/en/benefits.



USA, WA, Redmond - 143,700.00 - 194,400.00 USD annually"""

class JobD(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float
    education_requirements: list[str]
    responsibilities: list[str]

jobd_schema = JobD.model_json_schema()

system_prompt = f"""Your are an expert HR assistant. Your job is to analyze job descriptions and extract relevant JSON information based on the provided schema. The schema is as follows: {jobd_schema}. Please ensure that the extracted information is accurate and complete, adhering to the schema's structure and data types. If any information is missing or cannot be determined from the job description, please use null or an empty list as appropriate. The job description to analyze is provided below. Please return the extracted information in valid JSON format without any additional commentary or explanation."""

user_prompt = f"""Analyze the following Job Description: {job_description}"""

message_system={
    "role": "system",
    "content": system_prompt
}

message_user={
    "role": "user",
    "content": user_prompt
}

response_format={
    "type": "json_object"}

messages=[message_system, message_user]
response=client.chat.completions.create(model=model, messages=messages, response_format=response_format)

answer=response.choices[0].message.content

raw_json=answer

import json
job_data=json.loads(raw_json)
job=JobD(**job_data)

print(job.minimum_experience)
print(job.required_skills)



class MatchResult(BaseModel):
    score: float
    details: dict

class Experience(BaseModel):
    company: str | None=None
    role: str | None=None
    duration: str | None=None
    description: str | None=None
    skills_used: list[str] = []
    
class Resume(BaseModel):
    name: str | None=None
    email: str | None=None
    phone: str | None=None
    total_experience: float | None=None
    projects: list[str] = []
    experiences: list[Experience] = []
    skills: list[str] = []
    education: list[str] = []
    certifications: list[str] = []


resume_schema=Resume.model_json_schema()
def final_score(job,resume):
    match_schema=MatchResult.model_json_schema()
    prompt=f"""You are an expert HR assistant. Your task is to evaluate the compatibility between a job description and a candidate's resume. The job description is provided in the following JSON format: {job.model_dump_json()}. The candidate's resume is provided in the following JSON format: {resume.model_dump_json()}. Please assess the alignment of the candidate's qualifications, skills, and experiences with the requirements of the job description. Provide a final score between 0 and 100, where 0 indicates no match and 100 indicates a perfect match. Additionally, provide detailed feedback on areas of strength and areas for improvement in relation to the job requirements. Return your evaluation in valid JSON format adhering to the following schema: {match_schema}."""
    message={
        "role": "user",
        "content" : prompt
    }
    messages=[message]
    response_format={
        "type": "json_object"
    }
    response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)
def parse_resume(resume_text):
    system_prompt = f"""
    You are an expert resume parser.

    Extract information from the resume based on its meaning,
    not only based on exact section headings.

    Different resumes may use different headings.

    For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships

    These may all contain relevant experience.

    Skills may also appear in the skills section, work experience,
    internships or projects.

    Return ONLY valid JSON matching this schema:

    {resume_schema}

    Important rules:

    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    """
    user_prompt = f"""
    Parse the following resume:

    {resume_text}
    """
    message_system={
        "role" : "system",
        "content" : system_prompt
    }
    message_user={
        "role" : "user",
        "content" : user_prompt
    }
    messages=[message_system, message_user]
    response_format={
        "type": "json_object"
    }
    response=client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = Resume(**data)
    return resume


from pypdf import PdfReader
from docx import Document
def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None



# lets do it now
BASE_DIR = Path(__file__).resolve().parent
resume_folder = BASE_DIR / "resumes"
all_results=[]
for file_path in resume_folder.iterdir():
    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        continue
    print("\nProcessing:", file_path.name)
    resume_text = read_resume(file_path)
    parsed_resume=parse_resume(resume_text) # llm call1
    time.sleep(5)
    result = final_score(job, parsed_resume) #llm caLL2
    #score and details
    #acount chtgpt
    # request bhejna shhur krega millions
    #chattgot server jam ho jayega
    time.sleep(5)
    print("Score:", result.score)
    all_results.append({
        "name": parsed_resume.name,
        "score": result.score,
        "details": result.details
    })
all_results.sort(
    key=lambda candidate: candidate["score"],
    reverse=True
)
top_2 = all_results[:2]
worst_2 = all_results[-2:]


print("TOP 2 CANDIDATES")
for candidate in top_2:

    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )

    print(candidate["details"])

print("LOWEST 2 CANDIDATES")
for candidate in worst_2:

    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )
    print(candidate["details"])